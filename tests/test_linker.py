"""Offline tests for zotero_linker (fake Zotero client, synthetic .docx).
    python tests/test_linker.py"""
import json, os, sys, tempfile
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from docx.oxml.ns import qn
import zotero_linker as zl


def item(key, lastnames, year, title="T", creators=None):
    return zl.Item(key, f"http://zotero.org/users/999/items/{key}", {"id": key, "title": title}, title, lastnames, year)


LIB = [
    item("AAA", ["Nasim", "Lee", "Park"], "2026"),
    item("BBB", ["Lee", "Park"], "2024"),
    item("CCC", ["Moore"], "1999"),
    item("DDD", ["Müller"], "2021"),
    item("EEE", ["Garcia"], "2025"),          # in press
    item("FFF", ["Roe"], ""),                 # undated
    item("GGG", ["Roe"], "2010"),
]


class FakeClient:
    def search(self, q, limit=50):
        sur = q.split()[0].lower()
        return [i for i in LIB if any(zl.norm(n) == zl.norm(sur) for n in i.lastnames[:1])]


def resolver():
    return zl.Resolver(FakeClient(), False)


def test_norm_and_parse():
    assert zl.norm("Müller") == zl.norm("Muller") == "muller"
    assert zl.parse_part("Garcia, in press").year == "in press"          # was truncated to "in pres"
    assert zl.parse_part("Moore, forthcoming").year == "forthcoming"
    assert zl.parse_part("Smith, 2020a").year == "2020"
    assert zl.parse_part("Roe, n.d.").year == "n.d."


def test_fits():
    r = resolver()
    for text, key in [("Garcia, in press", "EEE"), ("Roe, n.d.", "FFF"), ("Muller, 2021", "DDD"),
                      ("Lee & Park, 2024", "BBB"), ("Nasim et al., 2026", "AAA")]:
        p = zl.parse_part(text)
        it = r.resolve(p)
        assert it is not None and it.key == key, (text, it)
    assert r.resolve(zl.parse_part("Lee, 2024")) is None                  # 2 authors cited as 1
    assert r.resolve(zl.parse_part("Roe, 2011")) is None


def test_to_item_creator_types():
    c = zl.ZoteroClient(False, "999", None, None, 23119)
    raw = {"key": "K1", "data": {"title": "T", "date": "2020", "creators": [
        {"creatorType": "author", "lastName": "Smith"}, {"creatorType": "translator", "lastName": "Tran"}]},
        "csljson": {}}
    assert c._to_item(raw).lastnames == ["Smith"]                         # translator must not count
    raw["data"]["creators"] = [{"creatorType": "editor", "lastName": "Ed"}]
    assert c._to_item(raw).lastnames == ["Ed"]


def test_non_overlapping():
    ev = [(5, 20, "a", "a"), (0, 3, "b", "b"), (10, 15, "c", "c"), (20, 25, "d", "d")]
    assert [e[2] for e in zl.non_overlapping(ev)] == ["b", "a", "d"]


def para_text(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def test_document():
    with tempfile.TemporaryDirectory() as tmp:
        src, dst = os.path.join(tmp, "in.docx"), os.path.join(tmp, "out.docx")
        d = Document()
        d.add_paragraph("A (Nasim et al., 2026; Lee & Park, 2024, p. 12) and Moore (1999) done.")
        p = d.add_paragraph(); p.add_run("Split (Lee & ").bold = True; p.add_run("Park, 2024) end.")
        d.add_paragraph("Placeholder [[CITE: Moore 1999]] and unknown (Nobody, 1999).")
        d.add_paragraph("Müller (2021) and Garcia (in press).")
        d.add_paragraph("Refs: [[bibliography]] after")
        d.save(src)
        before = [para_text(p) for p in Document(src).paragraphs]
        r = resolver()
        zl.process(src, dst, r, "apa", "en-US", False)
        out = Document(dst)
        xml = out.element.xml
        n_fields = xml.count("ZOTERO_ITEM CSL_CITATION")
        assert n_fields == 6, n_fields     # paren(2 sources) + Moore narrative | split | placeholder | Muller + Garcia narratives
        assert xml.count("ZOTERO_BIBL") == 1
        texts = [para_text(p) for p in out.paragraphs]
        assert texts[0] == before[0] and texts[1] == before[1], texts            # visible text unchanged
        assert "(Nobody, 1999)" in texts[2] and "Nobody" in "\n".join(r.unresolved)
        assert texts[4].startswith("Refs: ") and texts[4].endswith(" after"), texts[4]  # text around marker kept
        for it in out.element.body.iter(qn("w:instrText")):                       # every field carries valid CSL JSON
            if "ZOTERO_ITEM" in it.text:
                data = json.loads(it.text.split("CSL_CITATION", 1)[1])
                assert data["citationItems"] and all(c["uris"] for c in data["citationItems"])
        # run again on the output: existing fields must be left alone
        dst2 = os.path.join(tmp, "out2.docx")
        zl.process(dst, dst2, resolver(), None, "en-US", False)
        assert Document(dst2).element.xml.count("ZOTERO_ITEM CSL_CITATION") == n_fields


def test_uri_and_user_id_warning():
    import io, contextlib
    raw = {"key": "K1", "data": {"title": "T", "date": "2020", "creators": []}, "csljson": {}, "library": {"id": 1}}
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        c = zl.ZoteroClient(False, None, None, None, 23119)
    assert "--user-id" in err.getvalue(), "must warn when no user id"
    assert c._to_item(raw).uri == "http://zotero.org/users/local/zlinker0/items/K1"   # not users/1
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        c2 = zl.ZoteroClient(False, "999", None, None, 23119)
        zl.ZoteroClient(True, "999", None, "key", 23119)
    assert err.getvalue() == "", "no warning when the id is given"
    assert c2._to_item(raw).uri == "http://zotero.org/users/999/items/K1"


def test_existing_bibliography_marker():
    with tempfile.TemporaryDirectory() as tmp:
        src, dst = os.path.join(tmp, "in.docx"), os.path.join(tmp, "out.docx")
        d = Document()
        d.add_paragraph("Text (Moore, 1999).")
        d.add_paragraph("[[bibliography]]")                       # marker alone -> paragraph removed
        d.add_paragraph("Keep this [[Bibliography]] tail")         # marker + text -> text kept
        d.add_paragraph("Last [[bibliography]]")
        d.save(src)
        # first run: no bibliography yet -> the first marker becomes the field, later ones are dropped
        zl.process(src, dst, resolver(), "apa", "en-US", False)
        out = Document(dst)
        xml = out.element.xml
        assert xml.count("ZOTERO_BIBL") == 1
        texts = [para_text(p) for p in out.paragraphs]
        assert texts[-2:] == ["Keep this  tail", "Last "], texts          # marker removed, surrounding text kept
        assert "" not in texts[1:2] and len(texts) == 4, texts            # lone-marker paragraph became the bibliography field

        # document that already has a bibliography field: marker text removed, field untouched
        d = Document(dst)
        d.add_paragraph("Extra [[bibliography]] marker")
        d.add_paragraph("[[bibliography]]")
        src2, dst2 = os.path.join(tmp, "in2.docx"), os.path.join(tmp, "out2.docx")
        d.save(src2)
        zl.process(src2, dst2, resolver(), None, "en-US", False)
        out2 = Document(dst2)
        assert out2.element.xml.count("ZOTERO_BIBL") == 1
        assert all("[[" not in para_text(p) for p in out2.paragraphs)
        assert len(out2.paragraphs) == len(out.paragraphs) + 1     # "Extra ..." kept, lone marker paragraph gone


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn(); print(name, "OK")
