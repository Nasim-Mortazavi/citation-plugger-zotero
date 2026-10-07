"""Self-contained tests: build synthetic Zotero-style .docx files, run zotero_hyperlinks, check results.
    python tests/test_hyperlinks.py"""
import os, re, subprocess, sys, tempfile
from xml.sax.saxutils import escape
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls

HERE = os.path.join(os.path.dirname(__file__), "..")


def run_xml(text, rpr=""):
    return f'<w:r {nsdecls("w")}>{rpr}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def fld(kind, code, result_runs, ppr=""):
    f = lambda t: f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="{t}"/></w:r>'
    i = f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> {escape(code)} </w:instrText></w:r>'
    return [f("begin"), i, f("separate"), *result_runs, f("end")]


def build(path, paragraphs):
    """paragraphs: list of lists of raw run-xml strings"""
    d = Document()
    for runs in paragraphs:
        p = d.add_paragraph()
        for x in runs:
            p._p.append(parse_xml(x))
    d.save(path)


def cite(*texts, rpr=""):
    return fld("cite", "ADDIN ZOTERO_ITEM CSL_CITATION {}", [run_xml(t, rpr) for t in texts])


def bib(*entries):
    paras = []
    for k, e in enumerate(entries):
        runs = []
        if k == 0:
            runs += fld("bib", "ADDIN ZOTERO_BIBL {} CSL_BIBLIOGRAPHY", [])[:3]
        runs.append(run_xml(e))
        if k == len(entries) - 1:
            runs.append(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
        paras.append(runs)
    return paras


def link(src, *args):
    out = re.sub(r"(_linked)?\.docx$", "", src) + "_linked.docx"   # same rule as the script
    r = subprocess.run([sys.executable, os.path.join(HERE, "zotero_hyperlinks.py"), src, *args],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stdout + r.stderr
    return out, r.stdout


def links_of(path):
    d = Document(path)
    res = []
    for h in d.element.body.iter(qn("w:hyperlink")):
        res.append(("".join(t.text for t in h.iter(qn("w:t"))), h.get(qn("w:anchor"))))
    return res, "".join(t.text or "" for t in d.element.body.iter(qn("w:t")))


def verify(path):
    r = subprocess.run([sys.executable, os.path.join(HERE, "verify_links.py"), path],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stderr
    return r.stdout


def test_author_year(tmp):
    src = os.path.join(tmp, "ay.docx")
    paras = [
        [run_xml("Intro ")] + cite("(Smith, 2020, 2021; Lee & Park, 2024, p. 5)") + [run_xml(" end.")],
        [run_xml("Narrative: Nasim et al. ")] + cite("(2026, p. 3)") + [run_xml(" ok.")],
        [run_xml("Split ")] + cite("(Karlocai", " et al., 2014; ", "Moore, 1999)"),
        [run_xml("Missing ")] + cite("(Nobody, 1999)") + [run_xml(" and ")] + cite("(Li, 2020)"),
    ] + bib("Smith, J. (2020). A. J. 1.", "Smith, J. (2021). B. J. 2.", "Lee, A., & Park, B. (2024). C. J. 3.",
            "Nasim, N., et al. (2026). D. J. 4.", "Karlocai, K., et al. (2014). E. J. 5.",
            "Moore, M. (1999). F. J. 6.", "Liu, Q. (2020). G. J. 7.")
    build(src, paras)
    before = links_of(src)[1]
    out, stdout = link(src, "--blue")
    ls, after = links_of(out)
    assert after == before, f"text changed!\n{before}\n{after}"          # no duplicated/lost text
    got = {(t, a) for t, a in ls}
    for want in [("Smith, 2020", "_ZoteroRef_1"), ("2021", "_ZoteroRef_2"), ("Lee & Park, 2024", "_ZoteroRef_3"),
                 ("2026", "_ZoteroRef_4"), ("Moore, 1999", "_ZoteroRef_6")]:
        assert want in got, (want, sorted(got))
    assert any(a == "_ZoteroRef_5" for _, a in ls)                       # split-run citation linked
    assert not any(a == "_ZoteroRef_7" for _, a in ls), "(Li, 2020) must not link to Liu 2020"
    assert "Nobody 1999" in stdout and "Li 2020" in stdout, stdout
    v = verify(out); assert " 0 problems" in v, v
    out2, _ = link(out)                                                  # idempotent re-run
    assert links_of(out2) == links_of(out)


def test_numbered(tmp):
    src = os.path.join(tmp, "num.docx")
    sup = '<w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
    paras = [[run_xml("Claim")] + cite("1,3", rpr=sup) + [run_xml(". More")] + cite("9", rpr=sup)] + bib(
        "1. First ref", "2. Second ref", "3. Third ref")
    build(src, paras)
    out, stdout = link(src, "--blue")
    ls, _ = links_of(out)
    assert sorted(ls) == [("1", "_ZoteroRef_1"), ("3", "_ZoteroRef_3")], ls
    assert "9" in stdout                                                 # reported as unmatched
    d = Document(out)
    for r in d.element.body.iter(qn("w:hyperlink")):                     # rPr children in schema order
        tags = [c.tag.split("}")[1] for c in r.find(qn("w:r")).find(qn("w:rPr"))]
        assert tags == ["color", "u", "vertAlign"], tags
    v = verify(out); assert " 0 problems" in v, v


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as tmp:
        test_author_year(tmp); print("author-year OK")
        test_numbered(tmp); print("numbered OK")
