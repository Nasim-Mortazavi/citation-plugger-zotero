#!/usr/bin/env python3
"""
zotero_linker.py — turn plain-text citations in a Word thesis into LIVE Zotero
citation fields, exactly as if you had inserted them with the Zotero Word plugin.

    Before:   ...as shown recently (Nasim et al., 2026; Lee & Park, 2024, p. 12).
    After:    same text, but each citation is a real Zotero field. In Word,
              Zotero → Refresh re-formats them and fills the bibliography;
              Zotero → Edit Citation opens them like any other citation.

Recognised patterns
    (Nasim et al., 2026)   (Nasim, 2026)   (Nasim & Lee, 2026)   (Nasim and Lee 2026)
    (Nasim et al., 2026, p. 12)   (A, 2020; B et al., 2021)   (Nasim et al., 2026a)
    Nasim et al. (2026)   → narrative citation (author kept, year becomes the field)
    [[cite: free text query]]   → explicit query, if you prefer placeholders
    [[bibliography]]            → where the reference list goes (else appended)

Usage
    python zotero_linker.py thesis.docx                     # -> thesis_zotero.docx
    python zotero_linker.py thesis.docx --style ieee --interactive
    python zotero_linker.py thesis.docx --dry-run           # just report matches

Zotero connection: the running Zotero 7 desktop app (Settings → Advanced →
"Allow other applications on this computer to communicate with Zotero").
Tip: give --user-id (your numeric ID at zotero.org/settings/keys) so the fields
link directly to your library; without it Zotero will offer to re-link the
items the first time you press Refresh (the embedded data makes that painless).
"""
from __future__ import annotations

import argparse
import copy
import html
import json
import os
import random
import re
import string
import sys
import zipfile
from dataclasses import dataclass, field
from typing import Optional

import requests
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# ----------------------------------------------------------------------------- #
# Citation grammar
# ----------------------------------------------------------------------------- #
NAME = r"(?:(?:De|Van|Von|Del|Da|Di|Du|Le|La)\s)?[A-Z][\w'’\-]+(?:\s(?:van|von|de|der|del|da|di|le|la)\s[A-Z][\w'’\-]+)?"
AUTHORS = rf"(?:{NAME}(?:\set\s+al\.?|\s(?:&|and)\s{NAME}|(?:,\s{NAME})+,?\s(?:&|and)\s{NAME})?)"
YEAR = r"(?:19|20)\d{2}[a-z]?|n\.d\.|in press|forthcoming"
LOCATOR = r"(?:,?\s*(?:pp?\.|p|pp|chap\.?|ch\.?|§|para\.?)\s*[\d\-–,\s]+\d)?"
ONE_CITE = rf"(?:(?:see\s|cf\.\s|e\.g\.,?\s)?{AUTHORS},?\s(?:{YEAR}){LOCATOR})"
PAREN_RE = re.compile(rf"\(({ONE_CITE}(?:;\s*{ONE_CITE})*)\)")
NARRATIVE_RE = re.compile(rf"({AUTHORS})\s\(({YEAR}){LOCATOR}\)")
PLACEHOLDER_RE = re.compile(r"\[\[\s*cite\s*:\s*(.+?)\s*\]\]", re.IGNORECASE)
BIB_RE = re.compile(r"\[\[\s*bibliography\s*\]\]", re.IGNORECASE)
PART_RE = re.compile(rf"^(?:see\s|cf\.\s|e\.g\.,?\s)?({AUTHORS}),?\s({YEAR})(.*)$")
LOC_RE = re.compile(r"(pp?\.?|chap\.?|ch\.?|§|para\.?)\s*([\d\-–,\s]+\d)", re.IGNORECASE)


@dataclass
class Part:
    text: str                      # original text of this part
    surnames: list[str]
    etal: bool
    year: str
    locator: str = ""
    label: str = "page"
    prefix: str = ""               # "see ", "cf. "


def parse_part(s: str) -> Optional[Part]:
    m = PART_RE.match(s.strip())
    if not m:
        return None
    authors, year, tail = m.group(1), m.group(2), m.group(3)
    prefix = s.strip()[: m.start(1)].strip()
    etal = bool(re.search(r"\bet\s+al", authors))
    names = re.sub(r"\bet\s+al\.?", "", authors)
    surnames = [n.strip() for n in re.split(r"\s*(?:&|,|\band\b)\s*", names) if n.strip()]
    p = Part(s.strip(), surnames, etal, re.sub(r"[a-z]$", "", year), prefix=prefix)
    lm = LOC_RE.search(tail)
    if lm:
        lab = lm.group(1).lower()
        p.label = "chapter" if lab.startswith("ch") else "paragraph" if lab.startswith("para") else "section" if lab == "§" else "page"
        p.locator = lm.group(2).strip()
    return p


# ----------------------------------------------------------------------------- #
# Zotero
# ----------------------------------------------------------------------------- #
@dataclass
class Item:
    key: str
    uri: str
    csl: dict
    title: str
    lastnames: list[str]
    year: str

    def label(self):
        who = self.lastnames[0] + (" et al." if len(self.lastnames) > 2 else "") if self.lastnames else "Anon."
        return f"{who} ({self.year}) — {self.title[:70]}"


class ZoteroClient:
    def __init__(self, web: bool, user_id, group_id, api_key, port: int):
        self.web = web
        self.headers = {"Zotero-API-Version": "3"}
        self.user_id, self.group_id = user_id, group_id
        if web:
            if group_id:
                self.base = f"https://api.zotero.org/groups/{group_id}"
            elif user_id:
                self.base = f"https://api.zotero.org/users/{user_id}"
            else:
                sys.exit("--web needs --user-id or --group-id.")
            if not api_key:
                sys.exit("--web needs --api-key (or ZOTERO_API_KEY).")
            self.headers["Zotero-API-Key"] = api_key
        else:
            self.base = f"http://localhost:{port}/api/users/0"
        self.cache: dict[str, Item] = {}
        self.local_user_key: Optional[str] = None

    def _get(self, path, **params):
        params.setdefault("include", "data,csljson")
        try:
            r = requests.get(f"{self.base}{path}", params=params, headers=self.headers, timeout=20)
        except requests.ConnectionError:
            sys.exit("Could not reach Zotero. Is the desktop app running with the local API "
                     "enabled (Settings → Advanced)? Or use --web.")
        if r.status_code == 404 and not self.web:
            sys.exit("Zotero answered 404 — you need Zotero 7 with the local API enabled.")
        r.raise_for_status()
        d = r.json()
        return d if isinstance(d, list) else [d]

    def _to_item(self, raw: dict) -> Item:
        data, csl_raw = raw.get("data", {}), raw.get("csljson", {})
        if isinstance(csl_raw, str):
            csl_raw = json.loads(csl_raw)
        if isinstance(csl_raw, list):
            csl_raw = csl_raw[0] if csl_raw else {}
        csl = dict(csl_raw)
        lib = raw.get("library", {})
        lib_id = lib.get("id")
        if self.group_id:
            uri = f"http://zotero.org/groups/{self.group_id}/items/{raw['key']}"
        elif self.user_id:
            uri = f"http://zotero.org/users/{self.user_id}/items/{raw['key']}"
        elif lib_id:
            uri = f"http://zotero.org/users/{lib_id}/items/{raw['key']}"
        else:
            uri = f"http://zotero.org/users/local/{self.local_user_key or 'unknown'}/items/{raw['key']}"
        csl["id"] = raw["key"]
        last = [c.get("lastName") or c.get("name", "") for c in data.get("creators", [])
                if c.get("creatorType") in ("author", "editor", None) or True]
        m = re.search(r"\d{4}", data.get("date", ""))
        return Item(raw["key"], uri, csl, data.get("title", ""), last, m.group(0) if m else "")

    def search(self, query: str, limit=10) -> list[Item]:
        raws = self._get("/items", q=query, qmode="titleCreatorYear",
                         itemType="-attachment || note", limit=limit)
        items = [self._to_item(r) for r in raws]
        for it in items:
            self.cache[it.key] = it
        return items


def norm(s: str) -> str:
    return re.sub(r"[^a-z]", "", s.lower().replace("’", "'"))


# ----------------------------------------------------------------------------- #
# Matching
# ----------------------------------------------------------------------------- #
@dataclass
class Resolver:
    client: ZoteroClient
    interactive: bool
    used: dict[str, Item] = field(default_factory=dict)
    unresolved: list[str] = field(default_factory=list)
    log: list[str] = field(default_factory=list)
    decided: dict[str, Optional[Item]] = field(default_factory=dict)

    def resolve(self, p: Part) -> Optional[Item]:
        sig = f"{'|'.join(map(norm, p.surnames))}|{p.etal}|{p.year}"
        if sig in self.decided:
            return self.decided[sig]
        hits = self.client.search(f"{p.surnames[0]} {p.year}")
        good = [h for h in hits if self._fits(p, h)]
        item = self._pick(p, good, strict=True) if good else None
        self.decided[sig] = item
        if item:
            self.used[item.key] = item
            self.log.append(f"  ✔ {p.text:42} → {item.label()}")
        else:
            self.unresolved.append(p.text)
            self.log.append(f"  ✘ {p.text:42} → NOT FOUND")
        return item

    def resolve_query(self, q: str) -> Optional[Item]:
        hits = self.client.search(q)
        item = self._pick(Part(q, [], False, ""), hits, strict=False) if hits else None
        if item:
            self.used[item.key] = item
            self.log.append(f"  ✔ {q:42} → {item.label()}")
        else:
            self.unresolved.append(q)
            self.log.append(f"  ✘ {q:42} → NOT FOUND")
        return item

    @staticmethod
    def _fits(p: Part, it: Item) -> bool:
        if it.year != p.year and p.year not in ("n.d.", "in press", "forthcoming"):
            return False
        ln = [norm(x) for x in it.lastnames]
        want = [norm(x) for x in p.surnames]
        if not ln or not want or ln[0] != want[0]:
            return False
        if p.etal:
            return len(ln) >= 3
        if len(want) == 2:
            return len(ln) == 2 and ln[1] == want[1]
        return len(ln) == 1 or (len(want) > 2 and want == ln)

    def _pick(self, p: Part, hits: list[Item], strict: bool) -> Optional[Item]:
        if not hits:
            return None
        if len(hits) == 1:
            if not strict:
                self.log.append(f"  ! {p.text!r}: no exact author/year match, using closest hit")
            return hits[0]
        if not self.interactive:
            self.log.append(f"  ! {p.text!r}: {len(hits)} candidates, using the first (--interactive to choose)")
            return hits[0]
        print(f"\nWhich item is {p.text!r}?")
        for i, h in enumerate(hits, 1):
            print(f"  {i}. {h.label()}")
        while True:
            a = input("number (0 = skip): ").strip()
            if a.isdigit() and 0 <= int(a) <= len(hits):
                return hits[int(a) - 1] if int(a) else None


# ----------------------------------------------------------------------------- #
# Zotero field construction
# ----------------------------------------------------------------------------- #
def rand_id(n=8):
    return "".join(random.choices(string.ascii_letters + string.digits, k=n))


def citation_json(parts: list[tuple[Item, Part]], display: str, narrative=False) -> str:
    items = []
    for it, p in parts:
        ci = {"id": it.key, "uris": [it.uri], "itemData": it.csl}
        if p.locator:
            ci["locator"] = p.locator
            ci["label"] = p.label
        if p.prefix:
            ci["prefix"] = p.prefix
        if narrative:
            ci["suppress-author"] = True
        items.append(ci)
    return json.dumps({
        "citationID": rand_id(),
        "properties": {"formattedCitation": display, "plainCitation": display, "noteIndex": 0},
        "citationItems": items,
        "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json",
    }, ensure_ascii=False)


def _run(rpr, text=None, fld=None, instr=None):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    if fld:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), fld); r.append(e)
    if instr is not None:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = instr; r.append(e)
    if text is not None:
        e = OxmlElement("w:t"); e.set(qn("xml:space"), "preserve"); e.text = text; r.append(e)
    return r


def field_runs(rpr, code: str, display: str) -> list:
    return [_run(rpr, fld="begin"), _run(rpr, instr=f" {code} "),
            _run(rpr, fld="separate"), _run(rpr, text=display), _run(rpr, fld="end")]


def bib_field_runs(rpr):
    return field_runs(rpr, 'ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY',
                      "[Bibliography — press Zotero → Refresh in Word]")


# ----------------------------------------------------------------------------- #
# Document rewriting
# ----------------------------------------------------------------------------- #
def iter_paragraphs(doc):
    def walk(c):
        for p in c.paragraphs:
            yield p
        for t in c.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        yield from walk(s.header); yield from walk(s.footer)


def merge_runs(p: Paragraph):
    """Collapse all runs into the first one (used only for the [[bibliography]] marker)."""
    if len(p.runs) > 1:
        p.runs[0].text = p.text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)


def runs_in_fields(doc) -> set:
    """Every run that belongs to a Word field (existing Zotero citations, the
    bibliography, cross-references, ...). These are never read or rewritten."""
    inside, depth = set(), 0
    for r in doc.element.body.iter(qn("w:r")):
        fc = r.find(qn("w:fldChar"))
        kind = fc.get(qn("w:fldCharType")) if fc is not None else None
        if kind == "begin":
            depth += 1
        if depth or r.find(qn("w:instrText")) is not None:
            inside.add(r)
        if kind == "end":
            depth = max(0, depth - 1)
    return inside


def plain_groups(p: Paragraph, inside: set) -> list[list]:
    """Consecutive runs of the paragraph that are outside any field."""
    groups, cur = [], []
    for r in p.runs:
        if r._r in inside:
            if cur:
                groups.append(cur)
            cur = []
        else:
            cur.append(r)
    if cur:
        groups.append(cur)
    return groups


def merge_split_citations(group: list, pattern) -> None:
    """Where Word split a citation over several runs, merge just those runs
    (the rest of the paragraph and its formatting is left alone)."""
    text = "".join(r.text for r in group)
    bounds, pos = [], 0
    for r in group:
        bounds.append((pos, pos + len(r.text)))
        pos += len(r.text)
    for m in reversed(list(pattern.finditer(text))):
        idx = [i for i, (s, e) in enumerate(bounds) if s < m.end() and e > m.start()]
        if len(idx) > 1:
            group[idx[0]].text = "".join(group[i].text for i in idx)
            for i in idx[1:]:
                group[i]._r.getparent().remove(group[i]._r)


def rewrite_run(run, segments):
    """segments: list of ('text', str) | ('field', code, display). Replaces the run."""
    rpr = run._r.find(qn("w:rPr"))
    new = []
    for seg in segments:
        if seg[0] == "text":
            if seg[1]:
                new.append(_run(rpr, text=seg[1]))
        else:
            new.extend(field_runs(rpr, seg[1], seg[2]))
    for el in reversed(new):
        run._r.addnext(el)
    run._r.getparent().remove(run._r)


FIELD_PREFIX = "ADDIN ZOTERO_ITEM CSL_CITATION "


def find_events(text: str, resolver: Resolver, stats: dict) -> list[tuple[int, int, str, str]]:
    """Scan text for citation patterns; return sorted (start, end, field_code, display) events.

    start/end are character offsets into `text` — usable to locate the match in
    whatever the text came from (a docx run, a Word Selection, ...).
    """
    events = []   # (start, end, code, display)

    for m in PLACEHOLDER_RE.finditer(text):
        parts = []
        for q in m.group(1).split(";"):
            it = resolver.resolve_query(q.strip())
            if it:
                parts.append((it, Part(q, [], False, "")))
        if parts:
            disp = "(" + "; ".join(f"{it.lastnames[0] if it.lastnames else 'Anon.'}, {it.year}" for it, _ in parts) + ")"
            events.append((m.start(), m.end(), FIELD_PREFIX + citation_json(parts, disp), disp))

    for m in PAREN_RE.finditer(text):
        parts, ok = [], True
        for raw in m.group(1).split(";"):
            p = parse_part(raw)
            it = resolver.resolve(p) if p else None
            if it:
                parts.append((it, p))
            else:
                ok = False
        if ok and parts:
            events.append((m.start(), m.end(), FIELD_PREFIX + citation_json(parts, m.group(0)), m.group(0)))
        elif parts:
            stats["partial"] += 1
            resolver.log.append(f"  ! {m.group(0)}: only some sources found — left as text")

    for m in NARRATIVE_RE.finditer(text):
        tail = text[m.end(2):m.end() - 1]                # e.g. ", p. 3"
        p = parse_part(f"{m.group(1)}, {m.group(2)}{tail}")
        it = resolver.resolve(p) if p else None
        if it:
            disp = text[m.start(2) - 1:m.end()]          # "(2026, p. 3)"
            events.append((m.start(2) - 1, m.end(), FIELD_PREFIX + citation_json([(it, p)], disp, narrative=True), disp))

    events.sort()
    return events


def build_segments(text: str, resolver: Resolver, stats: dict) -> Optional[list]:
    """Scan text; return docx-run segments or None if nothing to change."""
    events = find_events(text, resolver, stats)
    if not events:
        return None
    segs, pos = [], 0
    for s, e, code, disp in events:
        if s < pos:
            continue
        segs.append(("text", text[pos:s]))
        segs.append(("field", code, disp))
        pos = e
        stats["fields"] += 1
    segs.append(("text", text[pos:]))
    return segs


def process(src, dst, resolver: Resolver, style: str, locale: str, dry_run: bool):
    doc = Document(src)
    stats = {"fields": 0, "partial": 0}
    bib_done = "ZOTERO_BIBL" in doc.element.xml      # document already has a Zotero bibliography
    if bib_done:
        print("Existing Zotero bibliography found — new references will be added to it.")
    any_pattern = re.compile("|".join(x.pattern for x in (PAREN_RE, NARRATIVE_RE, PLACEHOLDER_RE)))
    inside = runs_in_fields(doc)

    for p in iter_paragraphs(doc):
        if BIB_RE.search(p.text) and not bib_done:
            merge_runs(p)
            rewrite_run(p.runs[0], [("text", BIB_RE.split(p.text)[0])])
            for r in bib_field_runs(None):
                p._p.append(r)
            bib_done = True
            continue
        if not any_pattern.search(p.text):
            continue
        # citations split across runs (Word does this) → merge only those runs,
        # never touching existing fields
        for group in plain_groups(p, inside):
            merge_split_citations(group, any_pattern)
        for run in [r for r in p.runs if r._r not in inside]:
            segs = build_segments(run.text, resolver, stats)
            if segs:
                rewrite_run(run, segs)

    if not bib_done and resolver.used:
        h = doc.add_paragraph()
        try:
            h.style = "Heading 1"
        except KeyError:
            pass
        h.add_run("References")
        bp = doc.add_paragraph()
        for r in bib_field_runs(None):
            bp._p.append(r)

    print("Citations:\n" + "\n".join(resolver.log))
    print(f"\n{stats['fields']} citation(s) converted to Zotero fields, "
          f"{len(resolver.used)} unique reference(s).")
    if resolver.unresolved:
        print("\nNot found in Zotero (left as plain text):")
        for u in dict.fromkeys(resolver.unresolved):
            print(f"   - {u}")
    if dry_run:
        print("\n--dry-run: nothing written."); return
    doc.save(dst)
    existing = read_zotero_style(src)
    if style is None and existing:
        print(f"\nKeeping the document's Zotero style: {existing}")
    else:
        write_zotero_prefs(dst, style or "apa", locale)
    print(f"\nSaved → {dst}\nNext: open it in Word and click Zotero → Refresh.")


# ----------------------------------------------------------------------------- #
# Document preferences (so Zotero knows the style without asking)
# ----------------------------------------------------------------------------- #
def read_zotero_style(path: str) -> Optional[str]:
    """Style id already set in the document by the Zotero Word plugin, if any."""
    with zipfile.ZipFile(path) as z:
        if "docProps/custom.xml" not in z.namelist():
            return None
        custom = z.read("docProps/custom.xml").decode()
    chunks = re.findall(r'name="ZOTERO_PREF_(\d+)".*?<vt:lpwstr>(.*?)</vt:lpwstr>', custom, flags=re.S)
    data = html.unescape("".join(c for _, c in sorted(chunks, key=lambda x: int(x[0]))))
    m = re.search(r'<style id="([^"]+)"', data)
    return m.group(1) if m else None


def write_zotero_prefs(path: str, style: str, locale: str):
    style_uri = style if style.startswith("http") else f"http://www.zotero.org/styles/{style}"
    data = (f'<data data-version="3" zotero-version="7.0"><session id="{rand_id()}"/>'
            f'<style id="{style_uri}" locale="{locale}" hasBibliography="1" bibliographyStyleHasBeenSet="1"/>'
            f'<prefs><pref name="fieldType" value="Field"/>'
            f'<pref name="automaticJournalAbbreviations" value="true"/></prefs></data>')
    chunks = [data[i:i + 255] for i in range(0, len(data), 255)]

    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    custom = files.get("docProps/custom.xml", b"").decode() if "docProps/custom.xml" in files else ""
    if not custom:
        custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
                  'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"></Properties>')
        ct = files["[Content_Types].xml"].decode()
        if "docProps/custom.xml" not in ct:
            ct = ct.replace("</Types>", '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
            files["[Content_Types].xml"] = ct.encode()
        rels = files["_rels/.rels"].decode()
        if "docProps/custom.xml" not in rels:
            rels = rels.replace("</Relationships>", '<Relationship Id="rIdZoteroCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
            files["_rels/.rels"] = rels.encode()
    custom = re.sub(r'<property[^>]*name="ZOTERO_PREF_\d+".*?</property>', "", custom, flags=re.S)
    pids = [int(x) for x in re.findall(r'pid="(\d+)"', custom)]
    pid = max(pids, default=1) + 1
    props = ""
    for i, c in enumerate(chunks, 1):
        props += (f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{pid}" name="ZOTERO_PREF_{i}">'
                  f'<vt:lpwstr>{html.escape(c, quote=True)}</vt:lpwstr></property>')
        pid += 1
    files["docProps/custom.xml"] = custom.replace("</Properties>", props + "</Properties>").encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in files.items():
            z.writestr(n, b)


# ----------------------------------------------------------------------------- #
def main():
    for stream in (sys.stdout, sys.stderr):
        if stream.encoding and stream.encoding.lower() != "utf-8":
            stream.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx")
    ap.add_argument("-o", "--output")
    ap.add_argument("--style", default=None,
                    help="CSL style id (nature, apa, ieee, …). Default: keep the document's "
                         "existing Zotero style, else apa")
    ap.add_argument("--locale", default="en-US")
    ap.add_argument("--interactive", action="store_true", help="ask when several items could match")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--web", action="store_true", help="use api.zotero.org instead of the desktop app")
    ap.add_argument("--user-id", default=os.getenv("ZOTERO_USER_ID"), help="numeric zotero.org user ID (recommended)")
    ap.add_argument("--group-id", default=os.getenv("ZOTERO_GROUP_ID"))
    ap.add_argument("--api-key", default=os.getenv("ZOTERO_API_KEY"))
    ap.add_argument("--port", type=int, default=23119)
    a = ap.parse_args()
    out = a.output or re.sub(r"\.docx$", "", a.docx) + "_zotero.docx"
    client = ZoteroClient(a.web, a.user_id, a.group_id, a.api_key, a.port)
    process(a.docx, out, Resolver(client, a.interactive), a.style, a.locale, a.dry_run)


if __name__ == "__main__":
    main()
